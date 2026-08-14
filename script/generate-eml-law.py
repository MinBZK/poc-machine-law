#!/usr/bin/env python3
"""Genereer de EML-wet uit de officiele werkboeken van RVO.

De erkende maatregelenlijst is wetgeving: vier bijlagen bij de Omgevingsregeling
(BWBR0045528). RVO publiceert diezelfde bijlagen als werkboek, en dat werkboek is
de enige machineleesbare vorm ervan. Dit script leest de werkboeken en schrijft de
wet, zodat elke maatregeltekst letterlijk uit de bron komt en een nieuwe versie van
de lijst een hergeneratie is in plaats van handwerk.

    uv run script/generate-eml-law.py [--bron DIR] [--uit PAD]

Zonder --bron worden de werkboeken van rvo.nl gehaald.
"""

from __future__ import annotations

import argparse
import re
import sys
import urllib.request
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
import yaml

BASIS_URL = "https://www.rvo.nl/sites/default/files/2023-11"
WERKBOEKEN = {
    "algemeen": "erkende-maatregelenlijst-eml-2023-v-1-3.xlsx",
    "glastuinbouw": "erkende-maatregelenlijst-eml-2023-glastuinbouw-v-1-3.xlsx",
}

# De prefixletter van de maatregelcode bepaalt het onderdeel. In het algemene
# werkboek staat het onderdeel ook als kolom; die is voor GC4 leeg en in het
# glastuinbouwwerkboek helemaal. De afleiding uit de code is bij alle 148 gevulde
# rijen gelijk aan de kolom, en wordt daarom als gezaghebbend gebruikt.
ONDERDEEL_PER_PREFIX = {"F": "Faciliteiten", "G": "Gebouwen", "P": "Processen"}

# Artikel 4.14 Omgevingsregeling wijst de maatregelen aan voor milieubelastende
# activiteiten (F en P), artikel 5.29 die voor gebouwen (G). Beide artikelen kennen
# een afwijkende bijlage voor de glastuinbouwsector.
BIJLAGEN = {
    ("Faciliteiten", "algemeen"): "VII",
    ("Processen", "algemeen"): "VII",
    ("Gebouwen", "algemeen"): "XIV",
    ("Faciliteiten", "glastuinbouw"): "VIIaa",
    ("Processen", "glastuinbouw"): "VIIaa",
    ("Gebouwen", "glastuinbouw"): "XIVa",
}

GRONDSLAG_PER_ONDERDEEL = {
    "Faciliteiten": "artikel 5.15, vierde lid, Besluit activiteiten leefomgeving",
    "Processen": "artikel 5.15, vierde lid, Besluit activiteiten leefomgeving",
    "Gebouwen": "artikel 3.84, vijfde lid, Besluit bouwwerken leefomgeving",
}

OR = {
    "law": "Omgevingsregeling",
    "bwb_id": "BWBR0045528",
    "url": "https://wetten.overheid.nl/BWBR0045528/2024-01-01",
}
BAL = {
    "law": "Besluit activiteiten leefomgeving",
    "bwb_id": "BWBR0041330",
    "url": "https://wetten.overheid.nl/BWBR0041330/2024-01-01",
}
BBL = {
    "law": "Besluit bouwwerken leefomgeving",
    "bwb_id": "BWBR0041297",
    "url": "https://wetten.overheid.nl/BWBR0041297/2024-01-01",
}
WBM = {
    "law": "Wet belastingen op milieugrondslag",
    "bwb_id": "BWBR0007168",
    "url": "https://wetten.overheid.nl/BWBR0007168/2024-01-01",
}


# De ankers op wetten.overheid.nl dragen de hele hoofdstukindeling, niet alleen
# het artikelnummer. Een verzonnen "#Artikel3.205" bestaat niet en laat de lezer
# bovenaan een regeling van duizenden artikelen achter, precies waar een
# verwijzing hem vandaan moest houden. Overgenomen uit de gepubliceerde tekst.
ANKERS = {
    ("BWBR0045528", "4.14"): "Hoofdstuk4_Afdeling4.4_Artikel4.14",
    ("BWBR0045528", "4.14a"): "Hoofdstuk4_Afdeling4.4_Artikel4.14a",
    ("BWBR0045528", "5.29"): "Hoofdstuk5_Afdeling5.1_Paragraaf5.1.4_Artikel5.29",
    ("BWBR0045528", "5.30"): "Hoofdstuk5_Afdeling5.1_Paragraaf5.1.4_Artikel5.30",
    ("BWBR0041330", "5.15"): "Hoofdstuk5_Afdeling5.4_Paragraaf5.4.1_Artikel5.15",
    ("BWBR0041330", "3.205"): "Hoofdstuk3_Afdeling3.6_Paragraaf3.6.2_Artikel3.205",
    ("BWBR0041330", "3.211"): "Hoofdstuk3_Afdeling3.6_Paragraaf3.6.4_Artikel3.211",
    ("BWBR0041297", "3.84"): "Hoofdstuk3_Afdeling3.4_Paragraaf3.4.1_Artikel3.84",
    ("BWBR0007168", "60"): "HoofdstukVI_Afdeling4_Artikel60",
}


def grondslag(bron: dict, artikel: str, lid: str | None, uitleg: str) -> dict:
    """Bouw een legal_basis-blok conform legal_basis_format.md."""
    anker = ANKERS.get((bron["bwb_id"], artikel))
    if anker is None:
        raise KeyError(
            f"geen geverifieerd anker voor {bron['law']} artikel {artikel}; "
            "zoek het op in de gepubliceerde tekst in plaats van het te raden"
        )
    jc = f"jci1.3:c:{bron['bwb_id']}&artikel={artikel}"
    if lid:
        jc += f"&lid={lid}"
    blok = {
        "law": bron["law"],
        "bwb_id": bron["bwb_id"],
        "article": artikel,
        "url": f"{bron['url']}#{anker}",
        "juriconnect": f"{jc}&z=2024-01-01&g=2024-01-01",
        "explanation": uitleg,
    }
    if lid:
        blok["paragraph"] = lid
    return blok


def lees_werkboek(pad: Path) -> dict[str, list[dict[str, Any]]]:
    """Lees alle tabbladen als lijsten van rijen, met de kopregel als sleutels."""
    wb = openpyxl.load_workbook(pad, read_only=True, data_only=True)
    tabbladen = {}
    for naam in wb.sheetnames:
        rijen = list(wb[naam].iter_rows(values_only=True))
        # Het glastuinbouwwerkboek schrijft RVW_ID waar het algemene RVW ID schrijft.
        kop = [str(c).strip().replace("RVW_ID", "RVW ID") if c else f"kolom{i}" for i, c in enumerate(rijen[0])]
        tabbladen[naam] = [dict(zip(kop, rij, strict=False)) for rij in rijen[1:] if rij[0]]
    wb.close()
    return tabbladen


def tekst(waarde: Any) -> str | None:
    """Normaliseer een celwaarde tot nette tekst; leeg wordt None."""
    if waarde is None:
        return None
    # In het werkboek staan harde regeleindes als de letterlijke tekens _x000D_.
    schoon = str(waarde).replace("_x000D_", "\n").replace("\r\n", "\n").replace("\r", "\n")
    schoon = "\n".join(regel.rstrip() for regel in schoon.split("\n")).strip()
    return schoon or None


def _vergelijkbaar(waarde: str | None) -> str:
    """Vlak witruimte af, zodat twee schrijfwijzen van dezelfde zin gelijk worden."""
    return re.sub(r"\s+", " ", waarde or "").strip().lower()


def bouw_maatregelen(tabbladen: dict, lijst: str) -> list[dict]:
    """Zet een werkboek om in maatregel-objecten met hun randvoorwaarden."""
    technisch, economisch, alternatief = defaultdict(list), defaultdict(list), defaultdict(list)
    for rij in tabbladen["Technische rvw"]:
        technisch[rij["CodeMaat"]].append({"id": str(rij["RVW ID"]), "randvoorwaarde": tekst(rij["Randvoorwaarde"])})
    for rij in tabbladen["Economische rvw"]:
        economisch[rij["CodeMaat"]].append(
            {
                "id": str(rij["RVW ID"]),
                "realisatiemoment": tekst(rij["Realisatiemoment"]),
                "randvoorwaarde": tekst(rij["Randvoorwaarde"]),
            }
        )
    for rij in tabbladen["Alternatieve maatregelen"]:
        alternatief[rij["Maatregel code"]].append(
            {
                "naam": tekst(rij["Alternatieve Maatregel"]),
                "omschrijving": tekst(rij["Alternatieve Maatregel Omschrijving"]),
            }
        )

    maatregelen = []
    for rij in tabbladen["Maatregelen"]:
        code = rij["Maatregel Code"]
        onderdeel = ONDERDEEL_PER_PREFIX[code[0]]
        kolom_onderdeel = tekst(rij.get("Onderdeel"))
        if kolom_onderdeel and kolom_onderdeel != onderdeel:
            raise ValueError(f"{lijst} {code}: kolom zegt {kolom_onderdeel}, code zegt {onderdeel}")

        maatregel = {
            "code": code,
            "onderdeel": onderdeel,
            "categorie": tekst(rij["Categorie"]),
            "bijlage": BIJLAGEN[(onderdeel, lijst)],
            # Welk artikel deze maatregel aanwijst. De maatregelen voor gebouwen
            # zijn die bedoeld in artikel 3.84, vijfde lid, Bbl; die voor
            # milieubelastende activiteiten die bedoeld in artikel 5.15, vierde
            # lid, Bal. Twee grondslagen dus, en de wet moet per maatregel de
            # juiste noemen in plaats van alles onder één artikel te hangen.
            "grondslag": GRONDSLAG_PER_ONDERDEEL[onderdeel],
            "naam": tekst(rij["Naam Maatregel"]),
            "omschrijving": tekst(rij["Omschrijving Maatregel"]),
            "uitgangssituatie": tekst(rij["Uitgangssituatie"]),
            "doelmatig_beheer_en_onderhoud": tekst(rij["Doelmatig beheer en onderhoud (DBO)"]),
            "technische_randvoorwaarden": technisch.get(code, []),
            "economische_randvoorwaarden": economisch.get(code, []),
            "alternatieve_maatregelen": alternatief.get(code, []),
        }

        # Het werkboek noemt de randvoorwaarden twee keer: uitgeschreven op het
        # maatregelblad en genormaliseerd op een eigen tabblad. Die twee moeten
        # hetzelfde zeggen, anders mag er geen wet uit rollen.
        #
        # De Ja/Nee-kolommen 'Technische Randvoorwaarde', 'Economische Randvoorwaarden'
        # en 'Zelfstandig moment' worden niet overgenomen. Ze zijn spreadsheet-gemak in
        # plaats van bijlagetekst, en in het glastuinbouwwerkboek staan ze op zes
        # plaatsen verkeerd (FB8, FI4, GE1, GE2 en GE3 tweemaal), telkens tegen beide
        # inhoudelijke kolommen in. Het realisatiemoment staat in de bijlage zelf bij de
        # economische randvoorwaarde, en daar staat het hier ook.
        for kolom, sleutel in (
            ("Technische rvw", "technische_randvoorwaarden"),
            ("Economische rvw", "economische_randvoorwaarden"),
        ):
            uitgeschreven = _vergelijkbaar(tekst(rij.get(kolom)))
            leeg = uitgeschreven in ("", "niet van toepassing.", "niet van toepassing")
            if leeg != (not maatregel[sleutel]):
                raise ValueError(
                    f"{lijst} {code}: maatregelblad '{kolom}' en tabblad zijn het oneens over of er "
                    f"randvoorwaarden zijn (blad {uitgeschreven[:40]!r}, tabblad {len(maatregel[sleutel])})"
                )
            for stuk in maatregel[sleutel]:
                if _vergelijkbaar(stuk["randvoorwaarde"]) not in uitgeschreven:
                    raise ValueError(
                        f"{lijst} {code}: randvoorwaarde {stuk['id']} van het tabblad staat niet in '{kolom}'"
                    )

        maatregelen.append(maatregel)
    return maatregelen


def bouw_categorieen(per_lijst: dict[str, list[dict]]) -> list[dict]:
    """Verzamel de categorieen waarin de maatregelen zijn ingedeeld."""
    gezien: dict[str, dict] = {}
    for lijst, maatregelen in per_lijst.items():
        for m in maatregelen:
            cat = gezien.setdefault(
                m["categorie"], {"categorie": m["categorie"], "onderdeel": m["onderdeel"], "lijsten": []}
            )
            if lijst not in cat["lijsten"]:
                cat["lijsten"].append(lijst)
    return [gezien[k] for k in sorted(gezien)]


def bouw_wet(per_lijst: dict[str, list[dict]]) -> dict:
    """Stel de volledige wet samen."""
    aanwijzing = (
        "Artikel 4.14, tweede lid, en artikel 5.29, tweede lid, van de Omgevingsregeling wijzen de "
        "glastuinbouwbijlagen aan bij een activiteit als bedoeld in artikel 3.205 van het Besluit "
        "activiteiten leefomgeving (telen van gewassen in kassen), of bij een activiteit als bedoeld in "
        "artikel 3.211 van dat besluit (telen van gewassen in een gebouw, anders dan een kas) waarbij "
        "gebruik wordt gemaakt van het tarief, bedoeld in artikel 60, eerste lid, van de Wet belastingen "
        "op milieugrondslag."
    )

    parameters = [
        {
            "name": "TEELT_GEWASSEN_IN_KAS",
            "description": "Teelt het bedrijf gewassen in kassen? Niet als dat alleen bij een huishouden of beroep aan huis gebeurt, voor educatieve doeleinden, bij een onderzoeksinstelling of op volkstuinen.",
            "type": "boolean",
            "required": True,
            "legal_basis": grondslag(
                BAL,
                "3.205",
                None,
                "Artikel 3.205 Bal wijst het telen van gewassen in kassen aan als milieubelastende activiteit; "
                "het derde lid zondert huishoudens, beroep aan huis, educatieve doeleinden, onderzoeksinstellingen "
                "en volkstuinen uit. De aanwijzing bepaalt of de glastuinbouwbijlagen gelden.",
            ),
        },
        {
            "name": "TEELT_GEWASSEN_IN_GEBOUW_GEEN_KAS",
            "description": "Teelt het bedrijf gewassen in een gebouw dat geen kas is? Niet als dat alleen bij een huishouden of beroep aan huis gebeurt, voor educatieve doeleinden of bij een onderzoeksinstelling.",
            "type": "boolean",
            "required": True,
            "legal_basis": grondslag(
                BAL,
                "3.211",
                None,
                "Artikel 3.211 Bal wijst het telen van gewassen in een gebouw, anders dan een kas, aan als "
                "milieubelastende activiteit; het derde lid zondert huishoudens, beroep aan huis, educatieve "
                "doelen en onderzoeksinstellingen uit. Deze aanwijzing leidt alleen samen met het verlaagde "
                "energiebelastingtarief tot de glastuinbouwbijlagen.",
            ),
        },
        {
            "name": "MAAKT_GEBRUIK_VAN_VERLAAGD_ENERGIEBELASTINGTARIEF",
            "description": "Maakt het bedrijf gebruik van het verlaagde energiebelastingtarief voor de glastuinbouw?",
            "type": "boolean",
            "required": True,
            "legal_basis": grondslag(
                WBM,
                "60",
                "1",
                "Artikel 60, eerste lid, Wet belastingen op milieugrondslag regelt het verlaagde tarief voor de "
                "glastuinbouw. Dat tarief verandert de terugverdientijd en is daarom in artikel 4.14, tweede lid, "
                "en artikel 5.29, tweede lid, van de Omgevingsregeling voorwaarde voor de glastuinbouwbijlagen.",
            ),
        },
        {
            "name": "AANWEZIGE_CATEGORIEEN",
            "description": "De categorieen uit de erkende maatregelenlijst die bij het bedrijf voorkomen, bijvoorbeeld Perslucht of Binnenverlichting.",
            "type": "array",
            "required": True,
            "legal_basis": grondslag(
                BAL,
                "5.15",
                "1",
                "Artikel 5.15, eerste lid, Bal verplicht tot maatregelen met een terugverdientijd van ten hoogste "
                "vijf jaar. Een maatregel kan alleen gelden als de installatie of het gebouwdeel waarop hij ziet "
                "aanwezig is; die feitelijke situatie geeft de ondernemer op.",
            ),
        },
    ]

    momentopname = {"type": "point_in_time", "reference": "$calculation_date"}
    output = [
        {
            "name": "is_glastuinbouwsector",
            "description": "Gelden voor dit bedrijf de bijlagen die specifiek voor de glastuinbouwsector zijn vastgesteld?",
            "type": "boolean",
            "temporal": momentopname,
            "citizen_relevance": "primary",
            "legal_basis": grondslag(OR, "4.14", "2", aanwijzing),
        },
        {
            "name": "bijlage_milieubelastende_activiteiten",
            "description": "De bijlage bij de Omgevingsregeling met de maatregelen voor milieubelastende activiteiten (onderdelen Faciliteiten en Processen).",
            "type": "string",
            "temporal": momentopname,
            "citizen_relevance": "secondary",
            "legal_basis": grondslag(
                OR,
                "4.14",
                None,
                "Artikel 4.14 Omgevingsregeling wijst bijlage VII aan als de maatregelen bedoeld in artikel 5.15, "
                "vierde lid, Bal, en in afwijking daarvan bijlage VIIaa voor de glastuinbouwsector.",
            ),
        },
        {
            "name": "bijlage_gebouwen",
            "description": "De bijlage bij de Omgevingsregeling met de maatregelen voor gebouwen (onderdeel Gebouwen).",
            "type": "string",
            "temporal": momentopname,
            "citizen_relevance": "secondary",
            "legal_basis": grondslag(
                BBL,
                "3.84",
                "5",
                "Artikel 3.84, vijfde lid, Bbl draagt op bij ministeriele regeling de maatregelen aan te wijzen die "
                "voor gebouwen een terugverdientijd van ten hoogste vijf jaar hebben. Artikel 5.29 Omgevingsregeling "
                "wijst daarvoor bijlage XIV aan, en in afwijking daarvan bijlage XIVa voor de glastuinbouwsector.",
            ),
        },
        {
            "name": "maatregelen",
            "description": "De erkende maatregelen uit de voor dit bedrijf geldende bijlagen, in de categorieen die bij het bedrijf voorkomen, met de randvoorwaarden waaronder elke maatregel geldt.",
            "type": "array",
            "temporal": momentopname,
            "citizen_relevance": "primary",
            "legal_basis": grondslag(
                BAL,
                "5.15",
                "4",
                "Artikel 5.15, vierde lid, Bal bepaalt dat bij ministeriele regeling maatregelen worden aangewezen "
                "die in ieder geval een terugverdientijd van ten hoogste vijf jaar hebben; artikel 3.84, vijfde lid, "
                "Bbl doet dat voor gebouwen. De aangewezen maatregelen gelden onder de randvoorwaarden die de "
                "bijlage per maatregel noemt.",
            ),
        },
        {
            "name": "aantal_maatregelen",
            "description": "Het aantal erkende maatregelen dat op het bedrijf van toepassing kan zijn.",
            "type": "number",
            "temporal": momentopname,
            "citizen_relevance": "secondary",
            "legal_basis": grondslag(
                BAL,
                "5.15",
                "4",
                "Telling van de op grond van artikel 5.15, vierde lid, Bal en artikel 3.84, vijfde lid, Bbl "
                "aangewezen maatregelen die in de aanwezige categorieen vallen.",
            ),
        },
    ]

    definities = {
        "MAATREGELEN_ALGEMEEN": per_lijst["algemeen"],
        "MAATREGELEN_GLASTUINBOUW": per_lijst["glastuinbouw"],
        "CATEGORIEEN": bouw_categorieen(per_lijst),
    }

    acties = [
        {
            "output": "is_glastuinbouwsector",
            "operation": "OR",
            "values": [
                "$TEELT_GEWASSEN_IN_KAS",
                {
                    "operation": "AND",
                    "values": [
                        "$TEELT_GEWASSEN_IN_GEBOUW_GEEN_KAS",
                        "$MAAKT_GEBRUIK_VAN_VERLAAGD_ENERGIEBELASTINGTARIEF",
                    ],
                },
            ],
            "legal_basis": grondslag(OR, "4.14", "2", aanwijzing),
        },
        {
            "output": "bijlage_milieubelastende_activiteiten",
            "operation": "IF",
            "conditions": [
                {"test": "$is_glastuinbouwsector", "then": "VIIaa"},
                {"else": "VII"},
            ],
            "legal_basis": grondslag(
                OR,
                "4.14",
                None,
                "Bijlage VII geldt, tenzij de aanwijzing van het tweede lid opgaat; dan geldt bijlage VIIaa.",
            ),
        },
        {
            "output": "bijlage_gebouwen",
            "operation": "IF",
            "conditions": [
                {"test": "$is_glastuinbouwsector", "then": "XIVa"},
                {"else": "XIV"},
            ],
            "legal_basis": grondslag(
                OR,
                "5.29",
                None,
                "Bijlage XIV geldt, tenzij de aanwijzing van het tweede lid opgaat; dan geldt bijlage XIVa.",
            ),
        },
        {
            "output": "maatregelen",
            "operation": "FOREACH",
            "subject": {
                "operation": "IF",
                "conditions": [
                    {"test": "$is_glastuinbouwsector", "then": "$MAATREGELEN_GLASTUINBOUW"},
                    {"else": "$MAATREGELEN_ALGEMEEN"},
                ],
            },
            "where": {"operation": "IN", "subject": "$current.categorie", "values": "$AANWEZIGE_CATEGORIEEN"},
            "value": "$current",
            "legal_basis": grondslag(
                BAL,
                "5.15",
                "4",
                "Uit de voor het bedrijf geldende bijlagen worden de maatregelen genomen in de categorieen die bij "
                "het bedrijf voorkomen. De randvoorwaarden blijven bij de maatregel staan: de bijlage wijst de "
                "maatregel aan onder die voorwaarden, en de wet stelt niet vast dat eraan voldaan is.",
            ),
        },
        {
            "output": "aantal_maatregelen",
            "operation": "LENGTH",
            "subject": "$maatregelen",
            "legal_basis": grondslag(
                BAL,
                "5.15",
                "4",
                "Telling van de aangewezen maatregelen die in de aanwezige categorieen vallen.",
            ),
        },
    ]

    aantallen = {k: len(v) for k, v in per_lijst.items()}
    return {
        "$id": "https://raw.githubusercontent.com/MinBZK/poc-machine-law/refs/heads/main/schema/v0.1.7/schema.json",
        "uuid": "7c2e8f4a-91d3-4e5b-8a6f-2b0c9d1e3f57",
        "name": "Erkende maatregelenlijst energiebesparing (EML 2023)",
        "law": "omgevingswet/energiebesparing/maatregelen",
        "law_type": "FORMELE_WET",
        "legal_character": "BESCHIKKING",
        "decision_type": "ANDERE_HANDELING",
        "discoverable": "BUSINESS",
        "requires_manual_approval": False,
        # Als date en niet als string: het schema accepteert beide, maar de
        # engine draait er `datetime.combine` op en weigert de wet bij een
        # string ("combine() argument 1 must be datetime.date, not str").
        "valid_from": date(2024, 1, 1),
        "service": "RVO",
        "description": (
            "Bepaling welke erkende maatregelen ter verduurzaming van het energiegebruik voor een bedrijf gelden. "
            f"De wet draagt alle {aantallen['algemeen']} maatregelen van de algemene bijlagen (VII en XIV) en alle "
            f"{aantallen['glastuinbouw']} maatregelen van de glastuinbouwbijlagen (VIIaa en XIVa), met per maatregel "
            "de uitgangssituatie, de technische en economische randvoorwaarden, het doelmatig beheer en onderhoud en "
            "de alternatieve maatregelen, letterlijk zoals vastgesteld. Welke bijlagen gelden volgt uit de aanwijzing "
            "in artikel 4.14, tweede lid, en artikel 5.29, tweede lid, van de Omgevingsregeling. De wet stelt niet "
            "vast dat aan de randvoorwaarden van een maatregel is voldaan; die geeft zij mee zodat de ondernemer ze "
            "kan nagaan.\n"
            "Gegenereerd uit de werkboeken van RVO met script/generate-eml-law.py."
        ),
        "legal_basis": grondslag(
            OR,
            "4.14",
            None,
            "Artikel 4.14 Omgevingsregeling wijst de maatregelen aan bedoeld in artikel 5.15, vierde lid, Bal, en "
            "artikel 5.29 die bedoeld in artikel 3.84, vijfde lid, Bbl. Samen vormen die aanwijzingen de erkende "
            "maatregelenlijst.",
        ),
        "references": [
            {
                "law": "Omgevingsregeling, bijlage VII",
                "article": "4.14, eerste lid",
                "url": "https://wetten.overheid.nl/BWBR0045528/2024-01-01#BijlageVII",
            },
            {
                "law": "Omgevingsregeling, bijlage VIIaa (glastuinbouwsector)",
                "article": "4.14, tweede lid",
                "url": "https://wetten.overheid.nl/BWBR0045528/2024-01-01#BijlageVIIaa",
            },
            {
                "law": "Besluit bouwwerken leefomgeving",
                "article": "3.84, vijfde lid",
                "url": "https://wetten.overheid.nl/BWBR0041297/2024-01-01#Hoofdstuk3_Afdeling3.4_Paragraaf3.4.1_Artikel3.84",
            },
            {
                "law": "Omgevingsregeling, bijlage XIV",
                "article": "5.29, eerste lid",
                "url": "https://wetten.overheid.nl/BWBR0045528/2024-01-01#BijlageXIV",
            },
            {
                "law": "Omgevingsregeling, bijlage XIVa (glastuinbouwsector)",
                "article": "5.29, tweede lid",
                "url": "https://wetten.overheid.nl/BWBR0045528/2024-01-01#BijlageXIVa",
            },
            {
                "law": "Omgevingsregeling, bijlage XV en XVa (rekenmethodiek terugverdientijd)",
                "article": "4.14a en 5.30",
                "url": "https://wetten.overheid.nl/BWBR0045528/2024-01-01#BijlageXV",
            },
        ],
        "properties": {"parameters": parameters, "output": output, "definitions": definities},
        "variables": [],
        "requirements": [],
        "actions": acties,
    }


class Dumper(yaml.SafeDumper):
    """Houdt de wet leesbaar: meerregelige wetteksten als letterlijk blok."""


def _tekst_representer(dumper: yaml.Dumper, data: str):
    stijl = "|" if "\n" in data else None
    return dumper.represent_scalar("tag:yaml.org,2002:str", data, style=stijl)


Dumper.add_representer(str, _tekst_representer)


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--bron", type=Path, help="map met de werkboeken; zonder dit worden ze opgehaald van rvo.nl")
    p.add_argument(
        "--uit",
        type=Path,
        default=Path("laws/omgevingswet/energiebesparing/maatregelen/RVO-2024-01-01.yaml"),
        help="pad van de te schrijven wet",
    )
    args = p.parse_args()

    per_lijst = {}
    for lijst, bestand in WERKBOEKEN.items():
        if args.bron:
            pad = args.bron / bestand
        else:
            pad = Path(bestand)
            print(f"ophalen {BASIS_URL}/{bestand}", file=sys.stderr)
            urllib.request.urlretrieve(f"{BASIS_URL}/{bestand}", pad)  # noqa: S310
        per_lijst[lijst] = bouw_maatregelen(lees_werkboek(pad), lijst)
        print(f"{lijst}: {len(per_lijst[lijst])} maatregelen", file=sys.stderr)

    wet = bouw_wet(per_lijst)
    args.uit.parent.mkdir(parents=True, exist_ok=True)
    with args.uit.open("w", encoding="utf-8") as f:
        yaml.dump(wet, f, Dumper=Dumper, allow_unicode=True, sort_keys=False, width=100)
    print(f"geschreven: {args.uit}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
